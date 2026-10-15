"""Generate a French project presentation from the current AAPL artifacts."""

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
RESULTS = ROOT / "results"
ASSETS = RESULTS / "presentation_assets"
OUTPUT = RESULTS / "presentation_prevision_aapl.pptx"
NOTES_OUTPUT = RESULTS / "notes_orales_presentation.md"

BG = "F3F4EE"
PAPER = "FFFFFF"
INK = "17352D"
DARK = "10251F"
GREEN = "26705A"
MINT = "B9D4B1"
LIME = "D4E59A"
ORANGE = "E49A56"
RED = "C96652"
MUTED = "687971"
GRID = "DCE2DA"
FONT_HEAD = "Georgia"
FONT_BODY = "Aptos"


def color(value):
    return RGBColor.from_string(value)


def create_charts():
    ASSETS.mkdir(parents=True, exist_ok=True)
    features = pd.read_csv(DATA / "AAPL_features.csv", parse_dates=["Date"])
    backtest = pd.read_csv(RESULTS / "backtest_lstm_5d.csv", parse_dates=["date", "exit_date"])

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10,
            "axes.edgecolor": "#DCE2DA",
            "axes.labelcolor": "#687971",
            "xtick.color": "#687971",
            "ytick.color": "#687971",
            "text.color": "#17352D",
            "grid.color": "#DCE2DA",
            "grid.alpha": 0.6,
            "axes.spines.top": False,
            "axes.spines.right": False,
        }
    )

    fig, ax = plt.subplots(figsize=(8.0, 3.1), facecolor=f"#{DARK}")
    ax.set_facecolor(f"#{DARK}")
    ax.plot(features["Date"], features["Close"], color=f"#{LIME}", linewidth=2.2)
    ax.fill_between(features["Date"], features["Close"], color=f"#{GREEN}", alpha=0.25)
    ax.axis("off")
    fig.tight_layout(pad=0.1)
    fig.savefig(ASSETS / "cover_price.png", dpi=200, facecolor=fig.get_facecolor(), bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(9.0, 3.4))
    ax.plot(features["Date"], features["Close"], color=f"#{INK}", linewidth=1.5)
    ax.axvspan(pd.Timestamp("2015-01-01"), pd.Timestamp("2022-12-31"), color=f"#{GREEN}", alpha=0.10)
    ax.axvspan(pd.Timestamp("2023-01-01"), pd.Timestamp("2024-12-31"), color=f"#{ORANGE}", alpha=0.16)
    ax.axvspan(pd.Timestamp("2025-01-01"), features["Date"].max(), color=f"#{RED}", alpha=0.10)
    ax.set_ylabel("Clôture ($)")
    ax.set_xlim(features["Date"].min(), features["Date"].max())
    ax.xaxis.set_major_locator(mdates.YearLocator(2))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    ax.grid(axis="y")
    ax.legend(
        handles=[
            plt.Line2D([0], [0], color=f"#{GREEN}", linewidth=7, alpha=0.5, label="Train · 2015–2022"),
            plt.Line2D([0], [0], color=f"#{ORANGE}", linewidth=7, alpha=0.7, label="Validation · 2023–2024"),
            plt.Line2D([0], [0], color=f"#{RED}", linewidth=7, alpha=0.5, label="Test · 2025–2026"),
        ],
        frameon=False,
        ncol=3,
        loc="upper left",
        fontsize=8,
    )
    fig.tight_layout()
    fig.savefig(ASSETS / "price_split.png", dpi=190, facecolor="white", bbox_inches="tight")
    plt.close(fig)

    horizons = ["1 séance", "5 séances", "20 séances"]
    baseline_da = np.array([48.77, 46.05, 46.87])
    lstm_da = np.array([48.23, 51.23, 46.32])
    positions = np.arange(len(horizons))
    width = 0.32
    fig, ax = plt.subplots(figsize=(8.2, 4.0))
    ax.bar(positions - width / 2, baseline_da, width, color=f"#{MUTED}", label="Baseline persistance")
    ax.bar(positions + width / 2, lstm_da, width, color=f"#{GREEN}", label="LSTM")
    ax.axhline(50, color=f"#{ORANGE}", linestyle="--", linewidth=1.3, label="Repère 50 %")
    for index, value in enumerate(baseline_da):
        ax.text(index - width / 2, value + 0.35, f"{value:.1f}", ha="center", fontsize=9)
    for index, value in enumerate(lstm_da):
        ax.text(index + width / 2, value + 0.35, f"{value:.1f}", ha="center", fontsize=9, weight="bold")
    ax.set_xticks(positions, horizons)
    ax.set_ylim(40, 55)
    ax.set_ylabel("Directional Accuracy (%)")
    ax.grid(axis="y")
    ax.legend(frameon=False, ncol=3, loc="upper right", fontsize=8)
    fig.tight_layout()
    fig.savefig(ASSETS / "da_comparison.png", dpi=190, facecolor="white", bbox_inches="tight")
    plt.close(fig)

    tft_labels = ["TFT · 1 j", "TFT multi · 1 j", "TFT multi · 5 j", "TFT multi · 20 j"]
    tft_da = [48.50, 45.98, 43.10, 32.76]
    fig, ax = plt.subplots(figsize=(7.6, 3.8))
    y = np.arange(len(tft_labels))
    bar_colors = [f"#{GREEN}", f"#{ORANGE}", f"#{ORANGE}", f"#{RED}"]
    ax.barh(y, tft_da, color=bar_colors, height=0.58)
    for row, value in enumerate(tft_da):
        ax.text(value + 0.45, row, f"{value:.2f} %", va="center", fontsize=9, weight="bold")
    ax.set_yticks(y, tft_labels)
    ax.invert_yaxis()
    ax.set_xlim(0, 57)
    ax.axvline(50, color=f"#{MUTED}", linestyle="--", linewidth=1)
    ax.set_xlabel("Directional Accuracy (%)")
    ax.grid(axis="x")
    fig.tight_layout()
    fig.savefig(ASSETS / "tft_da.png", dpi=190, facecolor="white", bbox_inches="tight")
    plt.close(fig)

    backtest["drawdown"] = backtest["equity"] / backtest["equity"].cummax() - 1
    fig, axes = plt.subplots(2, 1, figsize=(8.5, 4.8), sharex=True, gridspec_kw={"height_ratios": [2, 1]})
    axes[0].plot(backtest["exit_date"], backtest["equity"], color=f"#{GREEN}", linewidth=2.0)
    axes[0].fill_between(backtest["exit_date"], 1, backtest["equity"], color=f"#{MINT}", alpha=0.45)
    axes[0].set_ylabel("Capital initial = 1")
    axes[0].grid(axis="y")
    axes[1].fill_between(backtest["exit_date"], backtest["drawdown"] * 100, 0, color=f"#{RED}", alpha=0.7)
    axes[1].set_ylabel("Drawdown (%)")
    axes[1].xaxis.set_major_locator(mdates.MonthLocator(interval=3))
    axes[1].xaxis.set_major_formatter(mdates.DateFormatter("%b %y"))
    axes[1].grid(axis="y")
    fig.tight_layout()
    fig.savefig(ASSETS / "backtest_equity.png", dpi=190, facecolor="white", bbox_inches="tight")
    plt.close(fig)

    return features, backtest


def add_text(slide, text, x, y, w, h, size=16, fill=INK, bold=False, font=FONT_BODY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Pt(0)
    frame.margin_right = Pt(0)
    frame.margin_top = Pt(0)
    frame.margin_bottom = Pt(0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(fill)
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = color(line)
    return shape


def add_line(slide, x1, y1, x2, y2, fill=GRID, width=1.0):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color(fill)
    line.line.width = Pt(width)
    return line


def add_image(slide, path, x, y, w, h):
    with Image.open(path) as image:
        image_ratio = image.width / image.height
    box_ratio = w / h
    if image_ratio > box_ratio:
        image_w = w
        image_h = w / image_ratio
        y += (h - image_h) / 2
    else:
        image_h = h
        image_w = h * image_ratio
        x += (w - image_w) / 2
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(image_w), height=Inches(image_h))


def add_header(slide, title, section, number):
    add_text(slide, section.upper(), 0.62, 0.34, 8.0, 0.22, 9, GREEN, True)
    add_text(slide, title, 0.62, 0.68, 12.0, 0.60, 26, INK, True, FONT_HEAD)
    add_line(slide, 0.62, 1.42, 12.72, 1.42, GRID, 1)
    add_line(slide, 0.62, 7.08, 12.72, 7.08, GRID, 0.8)
    add_text(slide, "Prévision financière multi-horizon · AAPL", 0.62, 7.17, 8.0, 0.18, 8, MUTED)
    add_text(slide, f"{number:02d} / 10", 11.78, 7.17, 0.94, 0.18, 8, MUTED, align=PP_ALIGN.RIGHT)


def new_slide(prs, title, section, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(BG)
    add_header(slide, title, section, number)
    return slide


def add_card(slide, x, y, w, h, title, body, accent=GREEN, title_size=15, body_size=12):
    add_rect(slide, x, y, w, h, PAPER)
    add_rect(slide, x, y, 0.07, h, accent)
    add_text(slide, title, x + 0.22, y + 0.19, w - 0.4, 0.35, title_size, INK, True)
    add_text(slide, body, x + 0.22, y + 0.63, w - 0.42, h - 0.78, body_size, MUTED)


def build_deck(features, backtest):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Prévoir AAPL : évaluation multi-horizon"
    prs.core_properties.subject = "LSTM, TFT et backtest sur rendements financiers"
    prs.core_properties.author = "Projet Prévision Financière Multi-Horizon"

    # 1 — Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(DARK)
    add_rect(slide, 0.64, 0.76, 0.12, 5.92, GREEN)
    add_text(slide, "PROJET DE MODÉLISATION · 2015–2026", 1.12, 0.93, 6.8, 0.3, 10, LIME, True)
    add_text(slide, "Prévoir AAPL\nà plusieurs horizons", 1.10, 1.62, 6.4, 1.78, 34, PAPER, True, FONT_HEAD)
    add_text(slide, "LSTM, TFT et simulation d’une stratégie à 5 séances", 1.12, 3.68, 5.9, 0.7, 17, "D7E3DC")
    add_line(slide, 1.12, 4.65, 6.45, 4.65, GREEN, 2.2)
    add_text(slide, "Question : le signal du modèle résiste-t-il à une baseline et aux coûts ?", 1.12, 4.93, 5.6, 0.85, 14, LIME)
    add_image(slide, ASSETS / "cover_price.png", 7.05, 1.55, 5.65, 4.25)
    add_text(slide, "AAPL  ·  rendements à 1, 5 et 20 séances", 7.12, 6.05, 5.55, 0.3, 10, "D7E3DC")
    add_text(slide, "SYNTHÈSE DES EXPÉRIMENTATIONS", 1.12, 6.82, 5.0, 0.25, 9, "B9C8BF", True)

    # 2 — Research question
    slide = new_slide(prs, "Le défi : prédire le signe, pas seulement lisser le prix", "01 · Problématique", 2)
    add_text(slide, "Un bon RMSE ne garantit pas une bonne décision de trading.", 0.72, 1.72, 11.7, 0.55, 20, INK, True, FONT_HEAD)
    horizon_items = [("01", "1 séance", "court terme"), ("05", "5 séances", "une semaine de marché"), ("20", "20 séances", "environ un mois")]
    for index, (number, label, sub) in enumerate(horizon_items):
        x = 0.72 + index * 4.12
        add_rect(slide, x, 2.65, 3.72, 1.38, PAPER)
        add_text(slide, number, x + 0.2, 2.84, 0.9, 0.62, 29, GREEN, True, FONT_HEAD)
        add_text(slide, label, x + 1.15, 2.87, 2.25, 0.37, 16, INK, True)
        add_text(slide, sub, x + 1.15, 3.32, 2.32, 0.31, 11, MUTED)
    add_card(slide, 0.72, 4.45, 5.75, 1.54, "La comparaison", "Baseline de persistance : le rendement récent sert de prédiction. Le modèle doit la dépasser hors échantillon.", GREEN)
    add_card(slide, 6.78, 4.45, 5.72, 1.54, "La métrique directionnelle", "DA = part des signes correctement prédits. 50 % est un repère, pas une preuve d’avantage statistique.", ORANGE)
    add_text(slide, "Cible : rendement logarithmique futur, pas prix brut", 0.74, 6.37, 11.5, 0.34, 14, GREEN, True)

    # 3 — Data and split
    slide = new_slide(prs, "Un split temporel pour tester sur le futur", "02 · Données", 3)
    add_image(slide, ASSETS / "price_split.png", 0.72, 1.68, 8.0, 3.12)
    add_card(slide, 9.02, 1.88, 3.45, 1.02, "TRAIN · 1 981", "2015–2022", GREEN, 14, 11)
    add_card(slide, 9.02, 3.05, 3.45, 1.02, "VALIDATION · 502", "2023–2024", ORANGE, 14, 11)
    add_card(slide, 9.02, 4.22, 3.45, 1.02, "TEST · 427", "2025–2026 · jusqu’au 16 septembre", RED, 14, 10)
    add_text(slide, "2 963 séances brutes → 2 910 lignes après indicateurs et cibles", 0.78, 5.35, 11.3, 0.38, 15, INK, True)
    add_text(slide, "Le test reste chronologiquement séparé; les 60 premiers jours de chaque split servent à former les fenêtres d’entrée.", 0.78, 5.86, 11.5, 0.63, 12, MUTED)

    # 4 — Pipeline
    slide = new_slide(prs, "Du cours historique à une décision simulée", "03 · Méthode", 4)
    steps = [
        ("01", "Cours AAPL", "OHLCV · yfinance"),
        ("02", "Features", "log-return · RSI · MACD"),
        ("03", "Fenêtre", "60 séances · 5 variables"),
        ("04", "Prévision", "rendements 1 / 5 / 20 j"),
        ("05", "Évaluation", "baseline · DA · backtest"),
    ]
    for index, (num, heading, detail) in enumerate(steps):
        x = 0.7 + index * 2.52
        add_rect(slide, x, 2.03, 2.18, 1.62, PAPER)
        add_text(slide, num, x + 0.18, 2.2, 0.5, 0.32, 11, ORANGE, True)
        add_text(slide, heading, x + 0.18, 2.63, 1.8, 0.36, 14, INK, True)
        add_text(slide, detail, x + 0.18, 3.08, 1.82, 0.44, 10, MUTED)
        if index < len(steps) - 1:
            add_line(slide, x + 2.18, 2.83, x + 2.47, 2.83, GREEN, 2)
    add_card(slide, 0.72, 4.22, 5.78, 1.55, "Prévention de la fuite", "StandardScaler ajusté uniquement sur le train; validation et test sont transformés avec ces paramètres. RSI ramené sur [0, 1].", GREEN)
    add_card(slide, 6.78, 4.22, 5.72, 1.55, "Choix de features", "Les bandes de Bollinger brutes sont écartées du LSTM : dépendantes du niveau de prix, elles dégradent la généralisation observée.", ORANGE)
    add_text(slide, "Outils : pandas · NumPy · ta · scikit-learn · PyTorch · PyTorch Forecasting / Lightning", 0.78, 6.24, 11.7, 0.4, 12, GREEN, True)

    # 5 — Models
    slide = new_slide(prs, "Deux familles de modèles, deux comportements", "04 · Modèles", 5)
    add_card(slide, 0.72, 1.82, 3.88, 3.55, "LSTM multi-horizon", "60 jours d’historique\n5 variables d’entrée\n32 unités récurrentes\n3 sorties : 1 / 5 / 20 j\n\nMSE · Adam · early stopping", GREEN, 17, 13)
    add_card(slide, 4.74, 1.82, 3.88, 3.55, "TFT · 1 jour", "Temporal Fusion Transformer\n18 473 paramètres\nQuantileLoss\nEntraînement 20 époques\n\nInterprétation des variables et de l’attention", ORANGE, 17, 13)
    add_card(slide, 8.76, 1.82, 3.82, 3.55, "TFT · 20 jours", "Même ordre de grandeur\nPrévision multi-pas en une passe\nÉvaluation agrégée à 1 / 5 / 20 j\n\nQuestion : le modèle apprend-il les horizons lointains ?", RED, 17, 13)
    add_text(slide, "Tous les modèles évalués sur la même séparation temporelle; les checkpoints relancés sont conservés séparément.", 0.76, 5.92, 11.7, 0.55, 13, MUTED)

    # 6 — LSTM results
    slide = new_slide(prs, "Le signal LSTM apparaît surtout à 5 séances", "05 · Résultats LSTM", 6)
    add_image(slide, ASSETS / "da_comparison.png", 0.67, 1.73, 7.7, 3.78)
    add_rect(slide, 8.75, 1.96, 3.78, 1.45, INK)
    add_text(slide, "+5,18 pts", 9.05, 2.18, 3.2, 0.6, 28, LIME, True, FONT_HEAD)
    add_text(slide, "d’écart de DA à 5 séances", 9.06, 2.85, 3.05, 0.31, 11, PAPER)
    add_card(slide, 8.75, 3.72, 3.78, 1.35, "Erreur à 5 séances", "RMSE LSTM 0,04360\ncontre 0,06694 pour la baseline", GREEN, 14, 12)
    add_text(slide, "À 1 et 20 séances, la DA du LSTM ne dépasse pas celle de la baseline.", 0.75, 5.96, 11.7, 0.55, 14, INK, True)
    add_text(slide, "Lecture : avantage local au seul horizon 5 jours, pas de supériorité générale.", 0.75, 6.48, 11.6, 0.3, 10, MUTED)

    # 7 — TFT results
    slide = new_slide(prs, "Le TFT réduit l’erreur, mais le multi-horizon se dégrade", "06 · Résultats TFT", 7)
    add_image(slide, ASSETS / "tft_da.png", 0.65, 1.78, 7.3, 3.65)
    add_card(slide, 8.25, 1.86, 4.18, 1.26, "TFT à 1 jour", "RMSE 0,01929 · MAE 0,01254\nDA 48,50 %", GREEN, 15, 12)
    add_card(slide, 8.25, 3.34, 4.18, 1.56, "TFT multi · DA", "1 jour : 45,98 %\n5 jours : 43,10 %\n20 jours : 32,76 %", RED, 15, 12)
    add_text(slide, "Le faible RMSE du TFT univarié ne se traduit pas par une meilleure direction. Le multi-pas est un résultat négatif à documenter.", 0.76, 5.83, 11.7, 0.75, 14, INK, True)

    # 8 — Backtest
    slide = new_slide(prs, "Backtest 5 jours : intéressant, encore exploratoire", "07 · Simulation", 8)
    add_image(slide, ASSETS / "backtest_equity.png", 0.62, 1.7, 7.75, 4.45)
    metrics = [("+50,65 %", "rendement cumulé", GREEN), ("1,471", "Sharpe annualisé", ORANGE), ("−10,17 %", "drawdown maximal", RED)]
    for index, (value, label, accent) in enumerate(metrics):
        y = 1.95 + index * 1.28
        add_rect(slide, 8.76, y, 3.74, 1.05, PAPER)
        add_rect(slide, 8.76, y, 0.07, 1.05, accent)
        add_text(slide, value, 9.0, y + 0.13, 3.18, 0.43, 22, INK, True, FONT_HEAD)
        add_text(slide, label, 9.0, y + 0.62, 3.15, 0.24, 10, MUTED)
    add_text(slide, "Long si la prédiction 5 j > 0, sinon cash · périodes non chevauchantes · 40 périodes investies sur 72", 0.74, 6.35, 11.8, 0.42, 11, MUTED)
    add_text(slide, "Coûts modélisés : 0,10 % de frais + 0,05 % de slippage par côté.", 0.74, 6.72, 11.8, 0.25, 10, GREEN, True)

    # 9 — Interpretation
    slide = new_slide(prs, "L’attention aide à lire le modèle, pas à prouver une cause", "08 · Interprétabilité", 9)
    add_image(slide, RESULTS / "tft_attention_encoder_variables.png", 0.68, 1.82, 6.3, 3.15)
    add_image(slide, RESULTS / "tft_attention_attention.png", 7.12, 1.82, 5.45, 3.15)
    add_text(slide, "Variables encodeur", 0.78, 5.08, 5.8, 0.3, 13, INK, True)
    add_text(slide, "Le RSI domine l’importance affichée (environ 39 %), devant MACD signal, MACD diff et log_return.", 0.78, 5.46, 5.82, 0.72, 11, MUTED)
    add_text(slide, "Poids d’attention moyens", 7.22, 5.08, 5.15, 0.3, 13, INK, True)
    add_text(slide, "Les poids varient selon les pas de l’encodeur. Une visualisation d’attention reste descriptive, pas une explication causale.", 7.22, 5.46, 5.2, 0.72, 11, MUTED)

    # 10 — Limits and takeaway
    slide = new_slide(prs, "Conclusion : un avantage ciblé, plusieurs questions ouvertes", "09 · Bilan", 10)
    add_rect(slide, 0.72, 1.82, 7.38, 1.35, INK)
    add_text(slide, "Le meilleur résultat du projet : le LSTM à 5 séances.", 1.0, 2.06, 6.8, 0.46, 19, PAPER, True, FONT_HEAD)
    add_text(slide, "Il dépasse la baseline en DA et en RMSE sur ce test, mais pas aux autres horizons.", 1.0, 2.58, 6.8, 0.36, 12, LIME)
    add_text(slide, "Sentiment : couverture disponible", 0.78, 3.64, 7.1, 0.35, 14, INK, True)
    add_line(slide, 1.0, 4.45, 7.5, 4.45, GRID, 4)
    add_line(slide, 1.0, 4.45, 2.24, 4.45, ORANGE, 8)
    add_text(slide, "2015", 0.9, 4.72, 0.8, 0.25, 10, MUTED)
    add_text(slide, "2016-10", 1.88, 4.72, 1.3, 0.25, 10, ORANGE, True)
    add_text(slide, "2026", 7.02, 4.72, 0.8, 0.25, 10, MUTED, align=PP_ALIGN.RIGHT)
    add_text(slide, "542 dates, uniquement janvier 2015–octobre 2016; pas d’intégration ni d’évaluation sentiment dans le LSTM.", 0.78, 5.14, 7.18, 0.67, 11, MUTED)
    add_card(slide, 8.48, 1.84, 4.02, 1.15, "À confirmer", "Une seule action et une seule période de test.", ORANGE, 14, 11)
    add_card(slide, 8.48, 3.18, 4.02, 1.15, "À compléter", "Historique sentiment plus long; Informer non implémenté.", RED, 14, 11)
    add_card(slide, 8.48, 4.52, 4.02, 1.15, "À retenir", "Un résultat prometteur n’est pas encore une stratégie robuste.", GREEN, 14, 11)
    add_text(slide, "Merci · Questions", 0.78, 6.42, 7.0, 0.42, 22, GREEN, True, FONT_HEAD)

    prs.save(OUTPUT)


def write_notes():
    notes = """# Trame orale — Prévision financière multi-horizon AAPL

Durée visée : 8 à 10 minutes. Les métriques sont celles du dernier entraînement relancé sur les données récupérées le 14 octobre 2026.

## 1. Prévoir AAPL à plusieurs horizons
Présenter le sujet et annoncer le fil : mesurer la qualité directionnelle, puis vérifier si elle se traduit dans une simulation avec coûts.

## 2. Le défi
Une erreur moyenne faible ne signifie pas que le modèle sait prévoir la hausse ou la baisse. La DA mesure la proportion de signes corrects. La baseline de persistance sert de comparaison simple, sans prétendre être un modèle sophistiqué.

## 3. Données
Décrire les 2 910 observations exploitables et le split strict par dates. Le test commence après l’entraînement et la validation. Les 60 premières observations de chaque segment sont nécessaires pour créer les fenêtres.

## 4. Méthode
Expliquer la fuite de données : si le scaler apprend aussi sur le test, il connaît indirectement la distribution future. Ici, il est ajusté sur le train uniquement. Les Bollinger brutes sont retirées du LSTM à cause de leur dépendance au niveau de prix.

## 5. Modèles
Le LSTM lit 60 jours de cinq indicateurs et produit trois rendements. Le TFT fournit un modèle de comparaison interprétable, testé en sortie 1 jour et en sortie multi-pas.

## 6. Résultats LSTM
Souligner le gain à 5 jours : DA 51,23 % contre 46,05 %, RMSE 0,04360 contre 0,06694. Dire aussi que le LSTM ne gagne pas sur la DA aux horizons 1 et 20 jours.

## 7. Résultats TFT
Le TFT à 1 jour a un RMSE de 0,01929 mais une DA de 48,50 %. Le TFT multi-horizon se dégrade, surtout à 20 jours. Cela illustre pourquoi il faut regarder plusieurs métriques.

## 8. Backtest
Décrire une stratégie long/cash, décidée par le signe de la prévision à 5 jours et rééquilibrée par blocs de cinq séances. Les coûts et le slippage sont appliqués par côté. Le rendement de 50,65 % est historique et ne garantit rien pour la suite; insister sur le test unique et les hypothèses d’exécution.

## 9. Attention
Le graphe d’importance place le RSI en premier, autour de 39 %. L’attention montre une pondération apprise, pas une relation causale. Ne pas conclure qu’une variable cause le mouvement du cours.

## 10. Bilan
Le LSTM apporte un signal favorable sur l’horizon 5 jours, pas une supériorité générale. Le sentiment n’est disponible que pour 542 dates de 2015 à octobre 2016 et n’a pas été testé comme feature. Informer reste à implémenter. Terminer sur la nécessité d’un test plus large, de plusieurs actifs et de données sentiment mieux couvertes.
"""
    NOTES_OUTPUT.write_text(notes, encoding="utf-8")


def main():
    features, backtest = create_charts()
    build_deck(features, backtest)
    write_notes()
    print(f"Présentation créée : {OUTPUT}")
    print(f"Notes orales créées : {NOTES_OUTPUT}")
    print(f"Slides : 10 | période données : {features['Date'].min().date()} à {features['Date'].max().date()}")
    print(f"Périodes backtest : {len(backtest)}")


if __name__ == "__main__":
    main()